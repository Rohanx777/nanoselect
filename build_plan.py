from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT / 'Nanomaterial_Hybrid_Model_Build_Plan.docx'
doc = Document()
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(26)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(5)

def p(s, style=None):
    return doc.add_paragraph(s, style)
def h(s): doc.add_heading(s, 2)
def bullet(s): p(s, 'List Bullet')
def code(s):
    a=p(s)
    a.paragraph_format.space_after=Pt(7)
    for r in a.runs: r.font.name='Consolas'; r.font.size=Pt(9)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,x in zip(t.rows[0].cells,headers): c.text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row): c.text=str(x)
    for i,row in enumerate(t.rows):
        if widths:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
        trpr=row._tr.get_or_add_trPr()
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
        no=OxmlElement('w:cantSplit'); trpr.append(no)
        for c in row.cells:
            tcpr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1' if i==0 else ('F4F6F8' if i%2==0 else 'FFFFFF')); tcpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            tcpr.append(borders)
            for a in c.paragraphs:
                a.paragraph_format.space_after=Pt(2)
                for r in a.runs: r.font.size=Pt(9); r.bold=i==0
    p('')
def page(title):
    doc.add_page_break(); doc.add_heading(title,1)

p('Nanomaterial selection model build plan','Title')
p('A local hybrid demonstration using published experimental data','Subtitle')
p('Prepared for Rohan | IDP project | 25 October 2026')
p('Build a reproducible Python notebook that considers 15 nanomaterial families, ranks supported candidates for a named pollutant, and shows numerical ML predictions only where an evaluated model applies. Use existing internet datasets and research. No full application, cloud deployment, or new laboratory experiments are required.')
h('The deliverable after two or three days')
bullet('One notebook that runs from top to bottom on your PC, plus a simple command line fallback.')
bullet('A catalogue containing all 15 materials and a source backed evidence table for a small supported pollutant list.')
bullet('One evaluated ML prediction module if the data passes the quality checks. TiO2 is the first candidate; more models are optional.')
bullet('A deterministic screening and ranking procedure with explanations, evidence links, and explicit unknowns.')
bullet('Saved model files, metrics, three demonstration cases, and an offline copy of the required data.')
h('The claim to make in your presentation')
p('“Our prototype recommends the best supported candidate among 15 material families for specified water conditions. It combines published evidence with numerical predictions in the scenarios covered by our trained model.”')
p('Do not claim that all 15 have trained models, that rankings establish universal material superiority, or that the system certifies drinking water quality. Unsupported candidates remain visible with an insufficient evidence status.')
h('How to use this plan')
p('Follow sections 1 to 9 in order. Use the copyable prompts in sections 10 and 11 to work with Codex in bounded stages. Treat the five hour budget as 300 minutes of planned assisted work, not a promise about account quota or wall clock availability. Dataset access and cleaning are the main uncertainties.')

page('1 Scope and schedule')
h('Freeze these decisions before coding')
bullet('Run locally with Python, pandas, scikit learn, matplotlib, joblib, openpyxl, Jupyter and pytest. Use no GPU framework or web interface.')
bullet('Support one pollutant per request. Start with methylene blue for the ML path if the dataset supports it. Add lead and one other pollutant as evidence scenarios only if verified sources are easy to obtain.')
bullet('Do not support arbitrary pollutant names. Display the supported list and return unsupported for unknown requests.')
bullet('Assess all 15 materials. Rank candidates with enough evidence; show incompatible and insufficient evidence candidates separately.')
bullet('Do not perform a live internet search during the demo. Acquire, verify and cache sources beforehand.')
table(['Stage','Assisted minutes','Deliverable'],[
('A Setup and data audit','45','Project skeleton and usable dataset decision'),
('B Evidence and catalogue','60','All 15 entries and sourced candidate rules'),
('C Train and evaluate','60','Baseline comparison and model eligibility'),
('D Integrate and rank','55','Deterministic recommendation pipeline'),
('E Verify and rehearse','40','Notebook and three cases run offline'),
('Reserve','40','Access failures or one repair cycle'),
('Total','300','A runnable demonstration')],[1.6,1.1,4.0])
h('Two day sequence')
p('Day 1: complete stages A, B and C. Save a checkpoint after each stage. Manually read extracted source passages and inspect the cleaned dataset between assisted sessions. Day 2: complete stages D and E, then use the reserve only for blockers. If a third day is available, use it for rehearsal and report preparation, not new features.')
h('Stop rules')
p('Stop searching for another ML dataset after the first 45 minute audit unless the first source is unusable. Do not add a second model task until the first end to end demo works. If no model beats the baseline reliably, keep its metrics in the notebook and use the evidence route for recommendations. Report the ML attempt honestly rather than manufacturing a successful result.')

page('2 Set up the local project')
p('Work inside C:\\Users\\Rohan\\Desktop\\IDP 2026. Inspect any AGENTS.md and existing files before making changes. Create a new nanomaterial_hybrid folder without overwriting existing work.')
h('PowerShell commands for a normal local Python installation')
code('py -3 -m venv .venv\n.\\.venv\\Scripts\\python.exe -m pip install pandas numpy scikit-learn matplotlib joblib openpyxl jupyter nbconvert ipykernel pytest\n.\\.venv\\Scripts\\python.exe -m pip freeze > requirements-lock.txt\n.\\.venv\\Scripts\\python.exe -m jupyter lab')
p('Run these from the new project folder. No activation script is required. If py is unavailable, ask Codex to find a usable Python interpreter before installing anything. Keep the resolved interpreter path in README.md.')
h('Minimal folder structure')
code('nanomaterial_hybrid/\n  README.md\n  PROJECT_STATE.md\n  requirements-lock.txt\n  data/raw/                 original downloads\n  data/processed/           cleaned model data\n  data/materials.csv        all 15 families\n  data/evidence.csv         verified evidence records\n  data/sources.csv          URLs and download metadata\n  config/ranking.json       published scoring rules\n  src/data_audit.py\n  src/train.py\n  src/recommend.py\n  notebooks/demo.ipynb\n  models/                   pipeline and model card\n  outputs/                  metrics and demo tables\n  tests/test_recommend.py')
h('Keep the implementation small')
p('Expose one main function: recommend(scenario). It returns a table with all 15 materials, their assessment status, preference rank when justified, any applicable prediction, source identifiers and a short reason. The notebook calls the same function as the command line demo so logic is not duplicated.')
p('PROJECT_STATE.md records the latest completed stage, dataset locations, exact run commands, open issues and the next task. Update it after each stage so you can continue without repeating the whole conversation.')
h('Hardware expectations')
p('The detected Intel Core i7 13650HX is sufficient for this tabular workflow. A GPU is unnecessary. RAM was not verified; 8 GB should be workable and 16 GB comfortable. Start Random Forest with 200 trees, random_state=42 and n_jobs=2. Keep the parameter search small.')

page('3 Acquire and audit existing data')
h('Download these sources first')
p('[S1] TiO2 data folder: https://github.com/mahyavs/GNN_photodegradation/tree/main/data')
p('The folder includes curated_data.xlsx and original_data.xlsx. Inspect sheets, headers, units and target definitions before training. The repository uses a transformed photodegradation rate target; verify its exact logarithm base and units against the paper and supplied code. Do not assume this file contains surface area or particle size.')
p('[S2] Associated TiO2 paper: https://pmc.ncbi.nlm.nih.gov/articles/PMC12126497/')
p('[S3] Magnetite data and code: https://github.com/Ahmed-Data/mnp-adsorption-ml')
p('The magnetite repository includes MNPs Data.xlsx and warns about incomplete descriptors. Its adsorption capacity target is not directly interchangeable with a photocatalysis rate constant. Use as an optional extension only after the TiO2 path works.')
h('Audit checklist before any training')
bullet('Save raw files unchanged. Record source URL, access date, repository commit or release, file hash and reuse terms. Public access does not by itself establish a reuse licence.')
bullet('List every sheet and column, row count, missing percentage, duplicate count, pollutant distribution and material coverage.')
bullet('Identify the exact target: percentage, capacity, rate constant or transformed rate. Record units and transformations explicitly.')
bullet('Locate paper identifiers, material batch identifiers and experiment identifiers. Trace back to supplements where practical.')
bullet('Exclude identifiers, measured final concentrations and any other target derived fields from input features.')
bullet('Separate pure materials from doped, coated and composite versions. Do not label a composite as the corresponding pure material.')
bullet('Choose inputs actually present in the usable dataset. Do not invent size, BET surface area, UV intensity or pollutant identity.')
h('Decision at the end of stage A')
p('Write outputs/data_audit.md with usable rows, independent studies where identifiable, features, target, units, allowed pollutant categories and limitations. Select one prediction task. If group provenance cannot be recovered in time, label a deduplicated random split as preliminary and explain that it may overestimate generalisation. Do not present it as independent study validation.')
p('Keep raw data, not just links, available for the final offline demonstration. Source acquisition may require manual downloads if a repository or supplement blocks automated access.')

page('4 Catalogue and evidence records')
p('Create these 15 catalogue entries immediately. Their presence means they are assessed, not that each is proven effective for each pollutant.')
table(['ID','Material family','Initial route to investigate'],[
('M01','Fe3O4 magnetite','Adsorption'),('M02','TiO2','Photocatalysis'),('M03','ZnO','Photocatalysis and adsorption'),('M04','Silver nanoparticles','Microbial inactivation'),('M05','nZVI','Reduction and adsorption'),('M06','Graphene oxide','Adsorption'),('M07','Reduced graphene oxide','Adsorption'),('M08','Carbon nanotubes','Adsorption'),('M09','Nano alumina','Adsorption'),('M10','MgO','Adsorption and precipitation'),('M11','MnO2','Adsorption and oxidation'),('M12','CeO2','Adsorption and catalysis'),('M13','Nano hydroxyapatite','Adsorption and ion exchange'),('M14','Nano clay montmorillonite','Adsorption and ion exchange'),('M15','Chitosan nanoparticles','Adsorption and complexation')],[.5,2.7,3.5])
p('These routes come from the supplied project list and are research leads. Verify each material and pollutant pairing before enabling a recommendation. Magnetic separation is a recovery step rather than the contaminant removal mechanism itself.')
h('Evidence table fields')
code('evidence_id, material_id, profile_id, pollutant_id, mechanism,\nwater_matrix, pH_min, pH_max, dose_g_L, time_min,\nlight_requirement, endpoint, result_value, result_unit,\nstudy_id, source_url, source_location, verified, limitations')
p('Use null for unknown values. A reported optimal pH is not automatically a validated operating range. Do not generalise from one tested point to every concentration or material grade.')
h('Research within the time limit')
p('Prioritise direct studies for the selected demo pollutants. Search by the specific material and pollutant, adding experimental terms such as pH, dose and supplementary data. A review can help find studies, but numerical rules should point back to the experiment. Cover all 15 in the catalogue; leave unverified pairs as insufficient evidence. Do not fabricate coverage to fill the table.')

page('5 Train one useful ML module')
h('Training sequence')
bullet('Load the audited dataset, normalise units and separate the inputs from the target. Keep the original target alongside any transformed version.')
bullet('Remove exact duplicates before splitting. Keep repeated measurements from one experiment, batch or paper together wherever identifiers allow.')
bullet('Reserve approximately 20 percent of study groups for final testing if the group count permits. Use grouped cross validation on the remaining groups to select a model. With too few groups, use leave one study out or report that evaluation is limited.')
bullet('Create a ColumnTransformer and Pipeline: numeric imputation and scaling for Ridge; categorical encoding with explicit unsupported category checks. Fit all preprocessing inside training folds.')
bullet('Compare DummyRegressor, Ridge and RandomForestRegressor. Limit Random Forest to a few choices of max_depth and min_samples_leaf. Do not run a large optimisation search.')
bullet('Choose the model on cross validation, then evaluate once on the untouched test groups. Save the pipeline, split identifiers and metrics.')
h('Metrics and acceptance')
p('Report MAE, RMSE and R squared, with the evaluated target and units in every plot or table. If training on a transformed target, report its metrics and, where mathematically valid, errors after inverse transformation. R squared can be negative. A negative value must not be hidden.')
p('Permit numerical recommendations only if the selected model improves on the baseline and its errors are useful for the demonstrated decisions. Do not promise a numerical accuracy threshold before examining the data. If validation is unstable, keep prediction experimental or disable it in final ranking.')
h('Prediction eligibility checks')
p('A saved model is used only for known pollutant categories, supported material profiles, complete required inputs and conditions within its declared domain. Check ranges and similarity to observed combinations; being inside every individual feature range is not enough to establish full coverage. Return evidence only when unsupported.')
h('Saved model card')
p('Record dataset version, input schema, target transformation, units, material and pollutant coverage, split strategy, metrics, baseline result, random seed, package versions and known limitations in models/model_card.json.')
p('References: [S4] https://scikit-learn.org/stable/common_pitfalls.html and [S5] https://scikit-learn.org/stable/modules/cross_validation.html. These document leakage prevention and group aware evaluation.')

page('6 Deterministic hybrid ranking')
p('The MVP uses deterministic rules for eligibility and a transparent multi criterion score for recommendation. Numerical ML output supplies additional evidence when its endpoint and conditions match the request. ML training may use randomness; fixed seeds and saved models make the demonstration reproducible in the pinned environment.')
h('Step one screen all candidates')
p('Assign supported, conditional, incompatible with setup, or insufficient evidence. A missing data entry means unknown, not zero effectiveness. Only rank supported candidates, or conditional candidates under explicitly accepted requirements. Put the other candidates in the same output table without a preference rank.')
h('Step two use a common decision rubric')
table(['Criterion','Weight','Rating anchors from zero to two'],[
('Treatment suitability','60%','0 evidence of failure for the request; 1 relevant potential with partial match; 2 relevant evidence supporting the objective'),
('Setup compatibility','25%','0 unavailable requirement; 1 conditional modification; 2 current setup meets requirements'),
('Recovery practicality','15%','0 documented impracticality; 1 additional documented step; 2 documented feasible recovery in the stated setup')],[1.3,.7,4.7])
p('These are project chosen preference weights and ordinal ratings, not learned scientific constants. Document the evidence for each assigned rating. Recovery of metal loaded particles still requires appropriate waste handling; easy collection is not proof of safe reuse.')
code('score = 100 * (0.60 * suitability/2\n             + 0.25 * setup/2\n             + 0.15 * recovery/2)')
p('For unknown criteria, calculate score_min and score_max by substituting zero and two respectively. Do not silently renormalise weights or impute the midpoint. Candidates with overlapping score ranges should be shown as tied or uncertain rather than forced into a unique best order.')
h('Step three attach ML without mixing units')
p('Show any valid numerical prediction in its own column. It may inform treatment suitability only when the requested endpoint and a predeclared target are comparable. Otherwise retain it as supporting information. Use numerical predictions to order materials only within a comparable validated subset, with the same endpoint and declared conditions. Never add a rate constant or removal percentage directly to the preference score.')
h('Step four check ranking stability')
p('Repeat rankings with weights 50/30/20 and 70/20/10. If the winner changes, label the recommendation sensitive to preferences. Use material_id only to keep table display order stable among ties, not to invent scientific superiority. Evidence strength is a separate descriptive field, never a claimed probability.')

page('7 Integrate the notebook demonstration')
h('Scenario schema')
code('scenario = {\n  "pollutant_id": "methylene_blue",\n  "concentration_mg_L": None,\n  "pH": None,\n  "time_min": None,\n  "dose_g_L": None,\n  "light_type": "unspecified",\n  "light_intensity": None,\n  "temperature_C": None,\n  "water_matrix": "laboratory_solution",\n  "available_equipment": [],\n  "requested_endpoint": "degradation_rate_constant"\n}')
p('Replace missing values only with actual user inputs or explicitly selected demonstration conditions from the training support. Display every assumed default. A scenario cannot produce numerical predictions if required values remain unknown.')
h('Pipeline behaviour')
code('validate request\nload catalogue and verified evidence\nfor each of the 15 materials:\n    check relevance and setup requirements\n    check whether a validated model applies\n    predict if eligible otherwise record a fallback reason\n    rate criteria with citations and explicit unknowns\ncalculate score ranges and supported preference tiers\nattach prediction units and source identifiers\nreturn all 15 assessments and a recommendation explanation')
h('Notebook cells in presentation order')
bullet('Project objective, scope and the list of 15 material families.')
bullet('Load data and display coverage, missingness and source records.')
bullet('Show saved model metrics and a predicted versus measured test plot.')
bullet('Set a scenario and run recommend(scenario).')
bullet('Display ranked candidates first, then all other assessments; show model coverage and evidence links.')
bullet('Change one condition and explain the resulting change or loss of prediction eligibility.')
bullet('Run the unsupported pollutant case and summarise limitations.')
p('Output columns: material_id, material_name, status, preference_tier, score_min, score_max, prediction, prediction_unit, prediction_basis, evidence_strength, source_ids, reason. Save tables as CSV and plots as PNG. These are evidence for your presentation, not substitutes for running the notebook.')

page('8 Verify and prepare the final demo')
h('Three required demonstration cases')
table(['Case','Input choice','Expected behaviour'],[
('ML plus evidence','A supported pollutant and conditions copied from the audited domain','Eligible ML prediction plus assessments of all 15 materials'),
('Changed setup','Same case with a required light condition removed or changed','Photocatalytic route becomes conditional or incompatible; numerical model falls back if needed'),
('Unsupported request','Pollutant absent from the supported database','Clear insufficient evidence response without fabricated predictions')],[1.1,2.6,3.0])
h('Meaningful automated checks')
bullet('Every recommendation returns exactly the 15 unique catalogue IDs.')
bullet('Changing input row order does not alter scores or preference tiers.')
bullet('Unknown material properties or evidence do not become zero removal or fabricated numerical predictions.')
bullet('A model refuses an unseen pollutant, missing required input and clearly out of domain conditions.')
bullet('Incompatible candidates cannot become the recommended winner.')
bullet('Scoring endpoints and units are not mixed; overlapping ranges remain uncertain.')
bullet('With the same saved model and inputs, repeated predictions are consistent within numerical tolerance.')
h('Reproduction commands')
code('.\\.venv\\Scripts\\python.exe -m pytest -q\n.\\.venv\\Scripts\\python.exe -m src.train\n.\\.venv\\Scripts\\python.exe -m src.recommend --scenario demo_ml\n.\\.venv\\Scripts\\python.exe -m jupyter nbconvert --to notebook --execute notebooks\\demo.ipynb --output demo_executed.ipynb')
p('Implement the module entry points to match these commands. Run them from the project root. Training is a development step; use the saved model during the live demonstration. Capture exact successful commands in README.md.')
h('Presentation sequence in five minutes')
p('Explain the material selection problem. Show catalogue and data provenance. Show honest held out model metrics. Run a supported case and explain the recommendation. Change the light setup to demonstrate deterministic screening. Finish with an unsupported case and the limits of the system.')
h('Completion checklist')
p('The notebook executes cleanly; raw files are cached; model metrics and units are visible; evidence claims have sources; all 15 materials appear; no fake data fills missing coverage; unknowns are explicit; the demo works without internet; PROJECT_STATE.md and README.md reflect the final implementation.')

page('9 Failure handling and claim boundaries')
table(['Problem','Action within the deadline'],[
('Download blocked','Try the associated paper supplement or one manual download. Preserve the source and file hash.'),
('Missing key descriptors','Use a defensible reduced feature set and state coverage limits. Do not invent material properties.'),
('No study identifiers','Recover provenance where practical. Otherwise mark random split evaluation preliminary and constrain claims.'),
('Model fails the baseline','Do not enable it for recommendation. Show the failed experiment and evidence only results.'),
('Not enough evidence for a material','Keep its catalogue row and return insufficient evidence for that scenario.'),
('No clear ranking winner','Return a tied shortlist and explain the missing information or preference sensitivity.'),
('Time nearly exhausted','Finish one notebook, one evaluated task and three cases. Remove optional model tasks and visual polish.')],[1.8,4.9])
h('Keep these distinctions clear')
p('Catalogue coverage: all 15 materials exist in the sample space. Evidence coverage: only verified material and pollutant pairs support recommendations. ML coverage: only the documented subset supports numerical predictions. Laboratory validation: not performed in this project. These are four separate claims.')
p('The final score expresses a declared recommendation policy. It is not experimentally measured removal efficiency. Literature evidence and ML predictions both carry limitations, and a trained model is not automatically stronger evidence than a relevant direct experiment.')
h('Optional additions only after completion')
bullet('A second material prediction task with compatible outcomes and sufficient data.')
bullet('Additional supported pollutants and better material profile descriptions.')
bullet('A compact local chart of recommendation score ranges.')
p('Exclude live web retrieval, neural networks, automated paper scraping, complex uncertainty calibration, multi pollutant optimisation and a full website from the deadline scope.')
h('Source register to maintain')
p('Record S1 and S2 for TiO2 data and methodology, S3 if magnetite is used, and S4 and S5 for evaluation methods. Add every material specific study as a new source entry with DOI or URL, access date, source location and the exact claim it supports. Treat the original supplied material list as an unverified catalogue seed, not experimental training evidence.')

page('10 Copyable Codex prompts for day one')
p('Paste one prompt at a time into the project chat. Ask Codex to complete the bounded stage, verify it and update PROJECT_STATE.md. Do not rerun completed stages after a quota reset or new chat.')
h('Prompt A Setup and data audit')
p('Work in C:\\Users\\Rohan\\Desktop\\IDP 2026. Read applicable AGENTS.md and inspect existing files. Create nanomaterial_hybrid without overwriting other work. Follow this Word plan. Build a local Python notebook project, not a full app. Acquire the existing TiO2 spreadsheets from https://github.com/mahyavs/GNN_photodegradation/tree/main/data and inspect the associated paper https://pmc.ncbi.nlm.nih.gov/articles/PMC12126497/. Save raw files and source metadata. Audit columns, units, missing values, duplicates, target transformation, study provenance and pollutant coverage. Do not invent data or assume surface area exists. Select one feasible prediction task and write outputs/data_audit.md. Stop this stage after the audit and skeleton; do not implement extra features. Update PROJECT_STATE.md with exact next steps.')
h('Prompt B Catalogue and evidence')
p('Continue from PROJECT_STATE.md. Create the 15 material catalogue from the plan with stable IDs. Research verified evidence for methylene blue and, only if time permits, lead and one additional pollutant. Prioritise primary experimental sources. Record formulation, pollutant, matrix, conditions, endpoint, units, study identifier, URL and source location. Preserve unknown values as null and unsupported pairs as insufficient evidence. Do not assume a composite result applies to a pure material or that a broad pH range follows from one optimum. Create evidence.csv, materials.csv and sources.csv. Keep the research within a 60 minute planning allocation and do not seek exhaustive coverage. Update the state file.')
h('Prompt C Train and evaluate')
p('Continue from the audited dataset. Implement leakage safe preprocessing and compare DummyRegressor, Ridge and a small Random Forest search with random_state=42. Split by source study or related experimental groups where recoverable. Keep final test groups untouched during model selection. Report MAE, RMSE and R squared with target units and transformations. Save the best useful pipeline, baseline comparison, split IDs, evaluation plot and model card. Build a prediction eligibility function rejecting unsupported pollutants, missing inputs and out of domain cases. If performance is not useful, disable numerical recommendations and preserve the honest results. Do not add deep learning or another dataset task. Update PROJECT_STATE.md.')

page('11 Copyable prompts and first action')
h('Prompt D Integrate deterministic ranking')
p('Continue from PROJECT_STATE.md. Implement recommend(scenario) using all 15 catalogue IDs, source backed eligibility rules, the 60/25/15 rubric in this plan, score bounds for unknown criteria and ties for overlapping bounds. Treat weights as a documented project policy. Check alternative weights 50/30/20 and 70/20/10. Keep evidence strength separate. Attach valid ML predictions in their own units; do not add numerical predictions to decision scores or automatically prioritise ML. Incompatible and insufficient evidence candidates remain visible but do not receive a preference rank. Build the minimal notebook and command line demo with one shared implementation. Save source IDs and explanations in each output. Update the state file.')
h('Prompt E Verify and rehearse')
p('Continue from PROJECT_STATE.md. Add meaningful checks for all 15 IDs, deterministic outputs, unsupported inputs, missing evidence, incompatible candidates, target units and ranking ties. Run tests and execute the notebook from top to bottom. Run three cases: an eligible ML scenario, a changed light setup and an unsupported pollutant. Cache all required data so the demonstration works offline. Repair failures within scope. Save outputs and provide exact commands in README.md, actual metrics and limitations. Do not add a website or claim lab validation. Update PROJECT_STATE.md with the finished status and a five minute presentation script.')
h('Prompt for resuming after an interruption')
p('Read PROJECT_STATE.md, README.md and outputs/data_audit.md in nanomaterial_hybrid. Inspect completed files before editing. Continue only the next unfinished stage in the Word build plan. Preserve verified data and completed work. Report any blocker that changes the feasibility of the demo. Do not restart dataset research, retrain completed models or add features without a concrete reason.')
h('Start here')
p('Open the IDP 2026 project in Codex and paste Prompt A. After its audit, check that a real target and usable inputs exist before moving to Prompt B. If the audit shows a critical data problem, resolve it within the reserve instead of committing to an unsupported model.')
p('The deadline target is a reproducible local demonstration with clear evidence and honest prediction coverage. Every candidate is assessed; a unique winner is produced only when the available evidence supports that decision.')

# Page numbers help navigation in an execution guide.
foot=sec.footer.paragraphs[0]
foot.alignment=2
r=foot.add_run('Build plan  |  '); r.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
doc.core_properties.title='Nanomaterial selection model build plan'
doc.core_properties.subject='Local hybrid recommendation demonstration in two or three days'
doc.core_properties.author='Rohan'
doc.save(OUT)
print(OUT)
